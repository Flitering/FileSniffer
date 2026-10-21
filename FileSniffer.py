"""
FileSniffer — анализ папок и архивов, поиск «неважных» данных и их безопасная очистка.

Запуск:
    python FileSniffer.py [папка]
"""
import os
import sys
import shutil
import zipfile
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from send2trash import send2trash

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QTableView, QVBoxLayout, QWidget, QPushButton,
    QMessageBox, QFileDialog, QLabel, QHBoxLayout, QHeaderView, QPlainTextEdit,
    QProgressBar, QLineEdit, QCheckBox, QSplitter, QAbstractItemView
)
from PyQt5.QtCore import (
    QSortFilterProxyModel, Qt, QTimer, QThread, pyqtSignal, QItemSelection, QItemSelectionModel
)
from PyQt5.QtGui import QStandardItemModel, QStandardItem, QColor, QBrush

# --- Настройки определения файлов ---
METADATA_EXTENSIONS = {'.tmp', '.db', '.sys', '.ini', '.log'}
NECESSARY_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.pdf', '.doc', '.docx', '.xls', '.xlsx'}
ARCHIVE_EXTENSIONS = {'.zip', '.rar', '.7z'}
PREVIEW_LIMIT = 10
# Запас свободного места при распаковке (доля от свободного места на диске)
FREE_SPACE_RESERVE = 0.9

STATUS_IMPORTANT = "Важное"
STATUS_NESTED = "Важное (во вложенных)"
STATUS_NORMAL = "Обычное"
STATUS_ERROR = "Ошибка"
TYPE_FOLDER = "Папка"
TYPE_ARCHIVE = "Архив"

COLUMNS = ["Путь к папке", "Количество файлов", "Файлы в папке (первые 10)", "Файлы метаданных",
           "Общий размер файлов", "Статус", "Тип"]
SORT_ROLE = Qt.UserRole
RECORD_ROLE = Qt.UserRole + 1


def _oem_encoding():
    """Кодировка имён в ZIP без флага UTF-8 (так пишут архиваторы Windows, напр. cp866)."""
    if sys.platform == "win32":
        try:
            import ctypes
            return f"cp{ctypes.windll.kernel32.GetOEMCP()}"
        except Exception:
            pass
    return "cp437"


OEM_ENCODING = _oem_encoding()


@dataclass
class Record:
    path: str
    kind: str
    file_count: int = 0
    files: list = field(default_factory=list)
    metadata_files: list = field(default_factory=list)
    size: int = 0
    important: bool = False
    nested_protected: bool = False
    encrypted: bool = False
    error: str = ""

    @property
    def protected(self):
        """Запись нельзя удалять: важное содержимое, шифрование, ошибка чтения или защищённые вложения."""
        return self.important or self.nested_protected or self.encrypted or bool(self.error)

    @property
    def status(self):
        if self.error:
            return f"{STATUS_ERROR}: {self.error}"
        if self.important:
            status = STATUS_IMPORTANT
        elif self.nested_protected:
            status = STATUS_NESTED
        else:
            status = STATUS_NORMAL
        if self.encrypted:
            status += " (Зашифрован)"
        return status

    @property
    def files_preview(self):
        return ', '.join(self.files[:PREVIEW_LIMIT]) + ('...' if self.file_count > PREVIEW_LIMIT else '')


def ext_of(name):
    return os.path.splitext(name)[1].lower()


def is_necessary_file(filename):
    return ext_of(filename) in NECESSARY_EXTENSIONS


def is_metadata_file(filename):
    return ext_of(filename) in METADATA_EXTENSIONS


def is_archive_file(filename):
    return ext_of(filename) in ARCHIVE_EXTENSIONS


def is_link(p):
    """Символическая ссылка или точка соединения (junction) — по ним не ходим."""
    isjunction = getattr(os.path, "isjunction", None)
    return os.path.islink(p) or bool(isjunction and isjunction(p))


def is_inside(child, parent):
    """child лежит строго внутри parent (с учётом регистра и нормализации путей)."""
    child = os.path.normcase(os.path.abspath(child))
    parent = os.path.normcase(os.path.abspath(parent))
    try:
        return child != parent and os.path.commonpath([child, parent]) == parent
    except ValueError:  # разные диски
        return False


def human_size(size):
    for unit in ("Б", "КБ", "МБ", "ГБ", "ТБ"):
        if size < 1024 or unit == "ТБ":
            return f"{size:.0f} {unit}" if unit == "Б" else f"{size:.1f} {unit}"
        size /= 1024


def zip_member_name(info):
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode(OEM_ENCODING)
    except UnicodeError:
        return info.filename


# --- Чтение архивов ---

def _list_zip(archive_path):
    """Возвращает [(имя, размер, является_папкой, является_ссылкой)], признак шифрования."""
    entries, encrypted = [], False
    with zipfile.ZipFile(archive_path, 'r') as z:
        for info in z.infolist():
            if info.flag_bits & 0x1:
                encrypted = True
            # Unix-симлинк в external_attr (S_IFLNK)
            link = (info.external_attr >> 16) & 0o170000 == 0o120000
            entries.append((zip_member_name(info), info.file_size, info.is_dir(), link))
    return entries, encrypted


def _list_rar(archive_path):
    import rarfile
    entries = []
    with rarfile.RarFile(archive_path) as rf:
        encrypted = rf.needs_password()
        for info in rf.infolist():
            if info.needs_password():
                encrypted = True
            entries.append((info.filename, info.file_size, info.is_dir(), info.is_symlink()))
    return entries, encrypted


def _list_7z(archive_path):
    import py7zr
    entries = []
    with py7zr.SevenZipFile(archive_path, mode='r') as archive:
        encrypted = archive.needs_password()
        if not encrypted:
            for info in archive.list():
                entries.append((info.filename, info.uncompressed or 0, info.is_directory,
                                getattr(info, "is_symlink", False)))
    return entries, encrypted


ARCHIVE_LISTERS = {'.zip': _list_zip, '.rar': _list_rar, '.7z': _list_7z}


def scan_archive(archive_path):
    record = Record(path=archive_path, kind=TYPE_ARCHIVE)
    lister = ARCHIVE_LISTERS.get(ext_of(archive_path))
    if lister is None:
        record.error = "неподдерживаемый формат"
        return record
    try:
        entries, record.encrypted = lister(archive_path)
    except ImportError as e:
        record.error = f"не установлена библиотека {e.name}"
        return record
    except Exception as e:
        record.error = str(e) or type(e).__name__
        return record
    for name, size, is_dir, _ in entries:
        if is_dir:
            continue
        record.file_count += 1
        record.size += size
        if len(record.files) < PREVIEW_LIMIT:
            record.files.append(name)
        if is_metadata_file(name):
            record.metadata_files.append(name)
        if is_necessary_file(name):
            record.important = True
    return record


# --- Сканирование ---

def scan_tree(root_path, on_progress=None, is_cancelled=lambda: False):
    """Обходит дерево и возвращает список Record для папок и архивов."""
    records = []
    folders = {}

    def on_error(err):
        rec = Record(path=err.filename or root_path, kind=TYPE_FOLDER, error=err.strerror or str(err))
        records.append(rec)
        folders[os.path.normcase(os.path.normpath(rec.path))] = rec

    for root, dirs, files in os.walk(root_path, onerror=on_error):
        if is_cancelled():
            return None
        dirs[:] = [d for d in dirs if not is_link(os.path.join(root, d))]
        rec = Record(path=root, kind=TYPE_FOLDER)
        archives = []
        for f in files:
            full = os.path.join(root, f)
            if is_archive_file(f):
                archives.append(full)
                continue
            rec.file_count += 1
            if len(rec.files) < PREVIEW_LIMIT:
                rec.files.append(f)
            if is_metadata_file(f):
                rec.metadata_files.append(f)
            if is_necessary_file(f):
                rec.important = True
            try:
                rec.size += os.path.getsize(full)
            except OSError:
                pass
        records.append(rec)
        folders[os.path.normcase(os.path.normpath(root))] = rec
        if on_progress:
            on_progress(root, len(records))
        for archive_path in archives:
            if is_cancelled():
                return None
            records.append(scan_archive(archive_path))

    _propagate_protection(records, folders, root_path)
    return records


def _propagate_protection(records, folders, root_path):
    """Помечает родительские папки, если внутри есть что-то защищённое."""
    root_key = os.path.normcase(os.path.normpath(root_path))
    for rec in records:
        if not (rec.important or rec.encrypted or rec.error or rec.nested_protected):
            continue
        parent = os.path.dirname(os.path.normpath(rec.path))
        while True:
            key = os.path.normcase(parent)
            folder = folders.get(key)
            if folder is not None:
                if folder.nested_protected:
                    break  # выше уже помечено
                folder.nested_protected = True
            if key == root_key or not is_inside(parent, root_path):
                break
            next_parent = os.path.dirname(parent)
            if next_parent == parent:
                break
            parent = next_parent


# --- Отчёт ---

def _cell_text(value):
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def save_report(records, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    headers = COLUMNS[:4] + ["Общий размер файлов (байт)"] + COLUMNS[5:]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row_idx, rec in enumerate(records, start=2):
        values = [rec.path, rec.file_count, rec.files_preview, ', '.join(rec.metadata_files),
                  rec.size, rec.status, rec.kind]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(value, int):
                cell.value = value
            else:
                # Явно строковый тип: имя файла вида "=HYPERLINK(...)" не станет формулой
                cell.value = _cell_text(value)
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFG", (80, 12, 60, 40, 18, 22, 10)):
        ws.column_dimensions[col].width = width
    wb.save(filename)


# --- Удаление ---

def check_deletable(target, root_path):
    """Повторная проверка непосредственно перед удалением. Возвращает (можно, причина)."""
    if not is_inside(target, root_path):
        return False, "вне сканируемой папки"
    if not os.path.lexists(target):
        return False, "не существует"
    if is_link(target):
        return False, "символическая ссылка"
    if os.path.isfile(target):
        if not is_archive_file(target):
            return False, "не архив"
        rec = scan_archive(target)
        return (False, rec.status) if rec.protected else (True, "")

    problems = []

    def on_error(err):
        problems.append(f"нет доступа к {err.filename}")

    for root, dirs, files in os.walk(target, onerror=on_error):
        dirs[:] = [d for d in dirs if not is_link(os.path.join(root, d))]
        for f in files:
            if is_necessary_file(f):
                return False, f"найден важный файл {os.path.join(root, f)}"
            if is_archive_file(f):
                rec = scan_archive(os.path.join(root, f))
                if rec.protected:
                    return False, f"архив {os.path.join(root, f)}: {rec.status}"
        if problems:
            return False, problems[0]
    if problems:
        return False, problems[0]
    return True, ""


def top_level_paths(paths):
    """Оставляет только самые верхние пути: удаление папки удаляет и всё вложенное."""
    result = []
    for p in sorted({os.path.normpath(p) for p in paths}, key=len):
        if not any(p == q or is_inside(p, q) for q in result):
            result.append(p)
    return result


def _remove_readonly(func, p, _exc):
    os.chmod(p, 0o666)
    func(p)


def delete_path(target, permanent):
    if not permanent:
        send2trash(target)
    elif os.path.isdir(target):
        shutil.rmtree(target, onerror=_remove_readonly)
    else:
        os.remove(target)


# --- Распаковка ---

def unique_path(p):
    if not os.path.exists(p):
        return p
    i = 1
    while os.path.exists(f"{p}_{i}"):
        i += 1
    return f"{p}_{i}"


def safe_member_target(dest_root, member_name):
    """Путь для элемента архива внутри dest_root; защита от Zip Slip (../, абсолютные пути, C:, ADS)."""
    name = member_name.replace("\\", "/")
    if ":" in name or name.startswith("/"):
        raise ValueError(f"недопустимое имя в архиве: {member_name}")
    target = os.path.normpath(os.path.join(dest_root, name))
    if not is_inside(target, dest_root):
        raise ValueError(f"элемент архива выходит за пределы папки: {member_name}")
    return target


def _check_entries(entries, encrypted, dest_root):
    if encrypted:
        raise ValueError("архив зашифрован")
    total = 0
    for name, size, _is_dir, link in entries:
        if link:
            raise ValueError(f"архив содержит символическую ссылку: {name}")
        safe_member_target(dest_root, name)
        total += size
    free = shutil.disk_usage(os.path.dirname(dest_root)).free
    if total > free * FREE_SPACE_RESERVE:
        raise ValueError(f"недостаточно места: нужно {human_size(total)}, свободно {human_size(free)}")


def extract_archive(archive_path, destination):
    ext = ext_of(archive_path)
    if ext not in ARCHIVE_LISTERS:
        raise ValueError("неподдерживаемый формат")
    dest_root = unique_path(os.path.join(destination, os.path.splitext(os.path.basename(archive_path))[0]))
    entries, encrypted = ARCHIVE_LISTERS[ext](archive_path)
    _check_entries(entries, encrypted, dest_root)
    os.makedirs(dest_root)
    if ext == '.zip':
        with zipfile.ZipFile(archive_path, 'r') as z:
            for info in z.infolist():
                target = safe_member_target(dest_root, zip_member_name(info))
                if info.is_dir():
                    os.makedirs(target, exist_ok=True)
                    continue
                os.makedirs(os.path.dirname(target), exist_ok=True)
                with z.open(info) as src, open(target, "wb") as dst:
                    shutil.copyfileobj(src, dst)
    elif ext == '.rar':
        import rarfile
        with rarfile.RarFile(archive_path) as rf:
            rf.extractall(dest_root)
    else:
        import py7zr
        with py7zr.SevenZipFile(archive_path, mode='r') as archive:
            archive.extractall(path=dest_root)
    return dest_root


def extract_item(item_path, destination):
    if os.path.isdir(item_path):
        dest = unique_path(os.path.join(destination, os.path.basename(os.path.normpath(item_path))))
        if is_inside(dest, item_path):
            raise ValueError("нельзя копировать папку внутрь самой себя")
        shutil.copytree(item_path, dest, symlinks=True)
        return dest
    return extract_archive(item_path, destination)


# --- Фоновые задачи ---

class ScanWorker(QThread):
    progress = pyqtSignal(str, int)
    done = pyqtSignal(object)

    def __init__(self, root_path):
        super().__init__()
        self.root_path = root_path
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        records = scan_tree(self.root_path, self.progress.emit, lambda: self._cancelled)
        self.done.emit(records)


class DeletionWorker(QThread):
    progress = pyqtSignal(int)
    logMessage = pyqtSignal(str)
    done = pyqtSignal(list)

    def __init__(self, paths, root_path, permanent=False):
        super().__init__()
        self.paths = paths
        self.root_path = root_path
        self.permanent = permanent

    def run(self):
        deleted = []
        total = len(self.paths)
        for i, p in enumerate(self.paths):
            try:
                ok, reason = check_deletable(p, self.root_path)
                if not ok:
                    self.logMessage.emit(f"Пропущено {p}: {reason}")
                else:
                    delete_path(p, self.permanent)
                    deleted.append(p)
                    action = "Удалено навсегда" if self.permanent else "Перемещено в корзину"
                    self.logMessage.emit(f"{action}: {p}")
            except Exception as e:
                self.logMessage.emit(f"Ошибка при удалении {p}: {e}")
            self.progress.emit(int((i + 1) / total * 100))
        self.done.emit(deleted)


class ExtractWorker(QThread):
    progress = pyqtSignal(int)
    logMessage = pyqtSignal(str)
    done = pyqtSignal(list)

    def __init__(self, paths, destination):
        super().__init__()
        self.paths = paths
        self.destination = destination

    def run(self):
        errors = []
        for i, p in enumerate(self.paths):
            try:
                dest = extract_item(p, self.destination)
                self.logMessage.emit(f"Готово: {p} → {dest}")
            except Exception as e:
                errors.append(p)
                self.logMessage.emit(f"Не удалось распаковать {p}: {e}")
            self.progress.emit(int((i + 1) / len(self.paths) * 100))
        self.done.emit(errors)


# --- Интерфейс ---

class RecordFilterProxy(QSortFilterProxyModel):
    def __init__(self):
        super().__init__()
        self.only_deletable = False

    def setOnlyDeletable(self, value):
        self.only_deletable = value
        self.invalidateFilter()

    def filterAcceptsRow(self, row, parent):
        if self.only_deletable:
            rec = self.sourceModel().item(row, 0).data(RECORD_ROLE)
            if rec.protected:
                return False
        return super().filterAcceptsRow(row, parent)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("FileSniffer — отчёт по папкам и архивам")
        self.resize(1300, 850)
        self.root_path = ""
        self.records = []
        self.worker = None

        self.model = QStandardItemModel()
        self.model.setHorizontalHeaderLabels(COLUMNS)
        self.proxy = RecordFilterProxy()
        self.proxy.setSourceModel(self.model)
        self.proxy.setSortRole(SORT_ROLE)
        self.proxy.setFilterCaseSensitivity(Qt.CaseInsensitive)
        self.proxy.setFilterKeyColumn(-1)
        self.initUI()
        self.setBusy(False)

    def initUI(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)

        # Верхняя панель
        top_layout = QHBoxLayout()
        self.choose_btn = QPushButton("Выбрать папку…")
        self.choose_btn.clicked.connect(self.chooseFolder)
        top_layout.addWidget(self.choose_btn)
        self.rescan_btn = QPushButton("Пересканировать")
        self.rescan_btn.clicked.connect(lambda: self.startScan(self.root_path))
        top_layout.addWidget(self.rescan_btn)
        self.folder_label = QLabel("Папка не выбрана")
        top_layout.addWidget(self.folder_label, 1)
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("Поиск по таблице…")
        self.search_edit.setClearButtonEnabled(True)
        self.search_edit.textChanged.connect(self.proxy.setFilterFixedString)
        top_layout.addWidget(self.search_edit, 1)
        self.only_deletable_cb = QCheckBox("Только неважные")
        self.only_deletable_cb.toggled.connect(self.proxy.setOnlyDeletable)
        top_layout.addWidget(self.only_deletable_cb)
        main_layout.addLayout(top_layout)

        # Таблица
        self.table = QTableView()
        self.table.setModel(self.proxy)
        self.table.setSortingEnabled(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setWordWrap(False)
        self.table.verticalHeader().setDefaultSectionSize(22)
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setStretchLastSection(True)
        for col, width in enumerate((520, 90, 300, 200, 110, 160, 60)):
            self.table.setColumnWidth(col, width)
        self.table.selectionModel().selectionChanged.connect(self.updateSummary)

        self.logTextEdit = QPlainTextEdit()
        self.logTextEdit.setReadOnly(True)
        self.logTextEdit.setMaximumBlockCount(10000)
        splitter = QSplitter(Qt.Vertical)
        splitter.addWidget(self.table)
        splitter.addWidget(self.logTextEdit)
        splitter.setStretchFactor(0, 4)
        splitter.setStretchFactor(1, 1)
        main_layout.addWidget(splitter, 1)

        # Панель действий
        action_layout = QHBoxLayout()
        self.select_non_important_btn = QPushButton("Выбрать все неважные")
        self.select_non_important_btn.clicked.connect(self.selectNonImportant)
        action_layout.addWidget(self.select_non_important_btn)
        self.delete_button = QPushButton("Удалить (в корзину)")
        self.delete_button.clicked.connect(lambda: self.deleteSelected(permanent=False))
        action_layout.addWidget(self.delete_button)
        self.perm_delete_button = QPushButton("Постоянное удаление")
        self.perm_delete_button.clicked.connect(lambda: self.deleteSelected(permanent=True))
        action_layout.addWidget(self.perm_delete_button)
        self.extract_button = QPushButton("Распаковать / скопировать…")
        self.extract_button.clicked.connect(self.extractSelected)
        action_layout.addWidget(self.extract_button)
        self.report_button = QPushButton("Сохранить отчёт Excel…")
        self.report_button.clicked.connect(self.saveReport)
        action_layout.addWidget(self.report_button)
        action_layout.addStretch()
        self.summary_label = QLabel()
        action_layout.addWidget(self.summary_label)
        main_layout.addLayout(action_layout)

        self.progressBar = QProgressBar()
        self.progressBar.setValue(0)
        main_layout.addWidget(self.progressBar)

    def setBusy(self, busy):
        has_data = bool(self.records)
        self.choose_btn.setEnabled(not busy)
        self.rescan_btn.setEnabled(not busy and bool(self.root_path))
        for btn in (self.select_non_important_btn, self.delete_button, self.perm_delete_button,
                    self.extract_button, self.report_button):
            btn.setEnabled(not busy and has_data)

    def appendLog(self, text):
        self.logTextEdit.appendPlainText(text)

    # --- Сканирование ---

    def chooseFolder(self):
        folder = QFileDialog.getExistingDirectory(self, "Выберите папку для сканирования", self.root_path)
        if folder:
            self.startScan(folder)

    def startScan(self, folder):
        if not folder or self.worker is not None:
            return
        self.root_path = os.path.normpath(folder)
        self.folder_label.setText(self.root_path)
        self.records = []
        self.model.removeRows(0, self.model.rowCount())
        self.setBusy(True)
        self.progressBar.setRange(0, 0)  # неопределённый прогресс
        self.appendLog(f"Начинаю сканирование: {self.root_path}")
        self.worker = ScanWorker(self.root_path)
        self.worker.progress.connect(self.onScanProgress)
        self.worker.done.connect(self.onScanFinished)
        self.worker.start()

    def onScanProgress(self, current, count):
        self.statusBar().showMessage(f"Просканировано {count}: {current}")

    def onScanFinished(self, records):
        self.worker = None
        self.progressBar.setRange(0, 100)
        self.progressBar.setValue(100)
        self.statusBar().clearMessage()
        self.records = records or []
        self.populate()
        deletable = sum(1 for r in self.records if not r.protected)
        self.appendLog(f"Сканирование завершено: {len(self.records)} записей, неважных — {deletable}.")
        self.setBusy(False)

    def populate(self):
        self.table.setSortingEnabled(False)
        self.model.removeRows(0, self.model.rowCount())
        protected_brush = QBrush(QColor(255, 235, 200))
        error_brush = QBrush(QColor(255, 210, 210))
        for rec in self.records:
            cells = [
                (rec.path, rec.path.lower()),
                (str(rec.file_count), rec.file_count),
                (rec.files_preview, rec.files_preview.lower()),
                (', '.join(rec.metadata_files), len(rec.metadata_files)),
                (human_size(rec.size), rec.size),
                (rec.status, rec.status),
                (rec.kind, rec.kind),
            ]
            row = []
            for text, sort_key in cells:
                item = QStandardItem(text)
                item.setEditable(False)
                item.setData(sort_key, SORT_ROLE)
                if rec.error:
                    item.setBackground(error_brush)
                elif rec.protected:
                    item.setBackground(protected_brush)
                row.append(item)
            row[0].setData(rec, RECORD_ROLE)
            row[0].setToolTip(rec.path)
            row[4].setToolTip(f"{rec.size:,} байт".replace(",", " "))
            self.model.appendRow(row)
        self.table.setSortingEnabled(True)
        self.updateSummary()

    # --- Выделение ---

    def recordAt(self, proxy_row):
        src = self.proxy.mapToSource(self.proxy.index(proxy_row, 0))
        return self.model.item(src.row(), 0).data(RECORD_ROLE)

    def selectedRecords(self):
        return [self.recordAt(i.row()) for i in self.table.selectionModel().selectedRows()]

    def updateSummary(self, *_):
        selected = self.selectedRecords()
        total = sum(r.size for r in selected)
        self.summary_label.setText(f"Выбрано: {len(selected)} ({human_size(total)}) | "
                                   f"Показано: {self.proxy.rowCount()} из {self.model.rowCount()}")

    def selectNonImportant(self):
        self.table.clearSelection()
        sel = QItemSelection()
        for row in range(self.proxy.rowCount()):
            if not self.recordAt(row).protected:
                sel.select(self.proxy.index(row, 0), self.proxy.index(row, len(COLUMNS) - 1))
        self.table.selectionModel().select(sel, QItemSelectionModel.Select | QItemSelectionModel.Rows)

    # --- Действия ---

    def deleteSelected(self, permanent):
        valid_paths = []
        for rec in self.selectedRecords():
            if os.path.normcase(os.path.normpath(rec.path)) == os.path.normcase(self.root_path):
                self.appendLog(f"Пропущено {rec.path}: родительская папка.")
            elif rec.protected:
                self.appendLog(f"Пропущено {rec.path}: {rec.status}.")
            else:
                valid_paths.append(rec.path)
        valid_paths = top_level_paths(valid_paths)
        if not valid_paths:
            QMessageBox.information(self, "Информация",
                                    "Нет элементов для удаления (важные, зашифрованные и нечитаемые элементы защищены).")
            return
        by_path = {os.path.normpath(r.path): r for r in self.records}
        total = sum(r.size for p, r in by_path.items() if any(p == v or is_inside(p, v) for v in valid_paths))
        if permanent:
            text = (f"Безвозвратно удалить {len(valid_paths)} элемент(ов), примерно {human_size(total)}?\n"
                    "Это действие нельзя отменить.")
        else:
            text = f"Переместить в корзину {len(valid_paths)} элемент(ов), примерно {human_size(total)}?"
        if QMessageBox.question(self, "Подтверждение", text, QMessageBox.Yes | QMessageBox.No,
                                QMessageBox.No) != QMessageBox.Yes:
            return
        self.setBusy(True)
        self.progressBar.setValue(0)
        self.appendLog("Начинается постоянное удаление..." if permanent else "Начинается удаление (в корзину)...")
        self.worker = DeletionWorker(valid_paths, self.root_path, permanent=permanent)
        self.worker.progress.connect(self.progressBar.setValue)
        self.worker.logMessage.connect(self.appendLog)
        self.worker.done.connect(self.onDeletionFinished)
        self.worker.start()

    def onDeletionFinished(self, deleted):
        self.worker = None
        if deleted:
            gone = lambda p: any(os.path.normpath(p) == d or is_inside(p, d) for d in deleted)
            self.records = [r for r in self.records if not gone(r.path)]
            for row in reversed(range(self.model.rowCount())):
                if gone(self.model.item(row, 0).data(RECORD_ROLE).path):
                    self.model.removeRow(row)
        self.updateSummary()
        self.setBusy(False)
        self.appendLog(f"Удаление завершено: удалено {len(deleted)}. "
                       "Размеры родительских папок обновятся после пересканирования.")
        QMessageBox.information(self, "Информация", f"Удаление завершено. Удалено элементов: {len(deleted)}.")

    def extractSelected(self):
        paths = [r.path for r in self.selectedRecords()]
        if not paths:
            QMessageBox.information(self, "Информация", "Пожалуйста, выберите элементы для распаковки.")
            return
        dest = QFileDialog.getExistingDirectory(self, "Выберите папку для распаковки")
        if not dest:
            return
        self.setBusy(True)
        self.progressBar.setValue(0)
        self.worker = ExtractWorker(paths, dest)
        self.worker.progress.connect(self.progressBar.setValue)
        self.worker.logMessage.connect(self.appendLog)
        self.worker.done.connect(self.onExtractFinished)
        self.worker.start()

    def onExtractFinished(self, errors):
        self.worker = None
        self.setBusy(False)
        if errors:
            QMessageBox.warning(self, "Ошибка", "Не удалось распаковать следующие элементы "
                                "(подробности в журнале):\n" + "\n".join(errors[:20]))
        else:
            QMessageBox.information(self, "Информация", "Распаковка завершена успешно.")

    def saveReport(self):
        name = os.path.basename(self.root_path.rstrip("\\/")) or "disk"
        default = os.path.join(os.path.expanduser("~"), "Documents", f"{name}_files_report.xlsx")
        filename, _ = QFileDialog.getSaveFileName(self, "Сохранить отчёт", default, "Excel (*.xlsx)")
        if not filename:
            return
        try:
            save_report(self.records, filename)
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Не удалось сохранить отчёт:\n{e}")
            return
        self.appendLog(f"Отчёт сохранён: {filename}")
        if QMessageBox.question(self, "Отчёт сохранён", "Открыть отчёт?",
                                QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
            try:
                os.startfile(filename)
            except Exception as e:
                self.appendLog(f"Не удалось открыть файл автоматически: {e}")

    def closeEvent(self, event):
        if self.worker is not None:
            if isinstance(self.worker, ScanWorker):
                self.worker.cancel()
                self.worker.wait()
            else:
                QMessageBox.information(self, "Подождите", "Дождитесь завершения текущей операции.")
                event.ignore()
                return
        event.accept()


def main():
    QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    if len(sys.argv) > 1 and os.path.isdir(sys.argv[1]):
        QTimer.singleShot(0, lambda: window.startScan(sys.argv[1]))
    else:
        QTimer.singleShot(0, window.chooseFolder)
    sys.exit(app.exec_())


if __name__ == '__main__':
    main()
